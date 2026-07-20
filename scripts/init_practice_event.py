#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
从考试实际数据回填 practice_event（学生做题历史）。

复用 init_from_exam.py 的解析逻辑，但产出【答题事件】（一次答题一行）。
【规范】kp_ids 存知识点 id（kp_id），通过 name→id 映射得到。

数据源：data/九年级数学考试实际数据/20251016-.../试题数据.xlsx + 小题得分表.xlsx
幂等：重跑先删 source='exam' 的旧记录再写。
"""
import os
import sys
import json

import openpyxl
import psycopg2.extras
from datetime import datetime

sys.path.insert(0, os.path.dirname(__file__))
from mastery_store import conn

ROOT = '/data/shanghui/Recommend_question'
EXAM = f'{ROOT}/data/九年级数学考试实际数据/20251016-0453-58eb-b45c-67523a043649'
QBANK = f'{ROOT}/data/初中_九年级_精简版.jsonl'

# 考试知识点名 → 标准 612 知识点名（与 init_from_exam 一致）
NORM = {
    '二次函数的定义': '二次函数的定义',
    '二次函数y=ax²+bx+c的图象和性质': '二次函数的图象与性质',
    '二次函数的图象和性质': '二次函数的图象与性质',
    'y=a（x-h）²+k的图象和性质': '二次函数的图象与性质',
    '二次函数的最值': '二次函数的最值',
    '二次函数图象与各项系数符号': '二次函数图象与系数a、b、c的关系',
    '二次函数图象的平移': '二次函数图象的平移',
    '二次函数与一元二次方程': '二次函数图象与一元二次方程的关系',
    '待定系数法求二次函数解析式': '用待定系数法确定二次函数表达式',
    '实际问题与二次函数': '二次函数的实际应用',
    '二次函数综合': '二次函数的综合应用',
    '反比例函数的定义': '反比例函数的定义',
    '反比例函数的性质': '反比例函数的图象与性质',
    '反比例函数与一次函数的综合': '反比例函数与一次函数的综合应用',
    '反比例函数、二次函数图象综合判断': '同一坐标系中函数图象的判断',
}

# ---- 0) name → kp_id 映射（从题库 kgPoints 反向）----
name2id = {}
with open(QBANK, encoding='utf-8') as f:
    for line in f:
        for kp in json.loads(line).get('kgPoints', []):
            kid, nm = kp.get('id'), kp.get('name')
            if kid and nm and nm not in name2id:
                name2id[nm] = kid

exam_ts = datetime(2025, 10, 16, 9, 0, 0)

# ---- 1) 试题数据：题号 → (标准KP名, 题型, 满分) ----
wb = openpyxl.load_workbook(f'{EXAM}/试题数据.xlsx', read_only=True, data_only=True)
qs = {}
for r in wb['题库数据'].iter_rows(values_only=True):
    if r[0] is None or not str(r[0]).isdigit():
        continue
    qno = int(r[0])
    ques_type = r[1]
    raw_kp = r[2]
    if raw_kp in NORM:
        qs[qno] = (NORM[raw_kp], ques_type, float(r[3]))
wb.close()

# ---- 2) 小题列 → (kp_id 列表, 该列满分, 题型)。题号11-14合并为"11-14"列 ----
col_map = {}
miss = set()
for qno, (kp_name, ques_type, full) in qs.items():
    key = '11-14' if 11 <= qno <= 14 else str(qno)
    col_map.setdefault(key, [[], 0.0, None])
    kid = name2id.get(kp_name)
    if kid:
        col_map[key][0].append(kid)
    else:
        miss.add(kp_name)
    col_map[key][1] += full
    col_map[key][2] = ques_type
if miss:
    print(f'警告: {len(miss)} 个标准知识点名在题库无对应 kp_id: {miss}')

# ---- 3) 解析学生得分 → 答题事件 ----
wb = openpyxl.load_workbook(f'{EXAM}/小题得分表.xlsx', read_only=True, data_only=True)
rows = list(wb['数学成绩表'].iter_rows(values_only=True))
wb.close()
hdr = [str(h) for h in rows[2][9:]]

events = []
for r in rows[4:]:
    sid = r[2]
    if sid is None or not str(sid).isdigit():
        continue
    sid = int(sid)
    subs = r[9:]
    for j, name in enumerate(hdr):
        if name not in col_map:
            continue
        kp_ids, full, ques_type = col_map[name]
        if not kp_ids:
            continue
        raw = subs[j] if j < len(subs) else None
        try:
            sc = float(raw) if raw not in (None, '') else 0.0
        except Exception:
            sc = 0.0
        if full <= 0:
            continue
        rate = sc / full
        events.append((
            sid, f'exam_20251016_{name}', kp_ids, ques_type,
            None, None, round(rate, 4), rate < 0.5, 'exam', exam_ts,
        ))

print(f'解析事件: {len(events)} 条（{len(set(e[0] for e in events))} 个学生）')

# ---- 4) 写库（幂等：先删旧 exam 记录）----
with conn() as cc, cc.cursor() as cur:
    cur.execute("DELETE FROM practice_event WHERE source = %s", ('exam',))
    psycopg2.extras.execute_values(cur,
        """INSERT INTO practice_event
           (student_id, question_id, kp_ids, ques_type, difficulty, d, score, is_wrong, source, ts)
           VALUES %s""", events)
    cc.commit()
print(f'已写入 {len(events)} 条')

# ---- 5) 抽查 ----
with conn() as cc, cc.cursor() as cur:
    cur.execute("SELECT count(*) FROM practice_event")
    print(f'表总数: {cur.fetchone()[0]}')
    cur.execute("SELECT count(DISTINCT kp) FROM practice_event, unnest(kp_ids) kp")
    print(f'涉及 kp_id 数:', cur.fetchone()[0])
print('完成。')
